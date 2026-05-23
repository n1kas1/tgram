"""Excel report generation for FundBot."""

from __future__ import annotations

from datetime import datetime
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font


def _fmt_date(value: Optional[datetime]) -> str:
    return value.strftime("%d.%m.%Y") if value else ""


def build_excel_report(path: str, campaign_rows: list[dict], user_rows: list[dict]) -> None:
    """Write a two-sheet Excel report to ``path``.

    ``campaign_rows`` — one dict per campaign; ``user_rows`` — one dict per
    participant with aggregate participation.
    """
    wb = Workbook()

    ws = wb.active
    ws.title = "Сборы"
    headers = [
        "Название", "Сумма", "На человека", "Собрано",
        "Подтверждено", "Ожидают", "Не оплатили", "Статус", "Создан", "Срок",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in campaign_rows:
        ws.append([
            r["title"], r["total"], r["per_user"], r["collected"],
            r["confirmed"], r["claimed"], r["unpaid"],
            "активен" if r["active"] else "закрыт",
            _fmt_date(r.get("created_at")), _fmt_date(r.get("due_date")),
        ])

    ws2 = wb.create_sheet("Участники")
    ws2.append(["Имя", "Участий в сборах", "Подтверждённых оплат"])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for r in user_rows:
        ws2.append([r["name"], r["campaigns"], r["confirmed"]])

    wb.save(path)
